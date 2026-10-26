"""Read-only identities and reader checks for the deterministic release build."""
import datetime as dt
import gzip
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT/'source/current-release.json.gz'
MANIFEST_PATH = ROOT/'source/current-release-manifest.json'
READER_PATH = ROOT/'source/Global_Border_Closures_Master_2016_2026.xlsx'
MODEL_PATH = ROOT/'source/current-reader-model.json'

def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def load_json(path):
    path=Path(path);raw=path.read_bytes()
    return json.loads(gzip.decompress(raw) if path.suffix=='.gz' else raw)

def native_sheets(release):
    """Reconstruct the recorded layout from authentic headers and preambles."""
    result={}
    for name,rows in release['tables'].items():
        headers=release['native_headers'][name]
        assert headers and len(headers)==len(set(headers)),('Invalid native headers',name)
        assert all(set(row)==set(headers) for row in rows),('Native row/header mismatch',name)
        result[name]=release['native_preambles'].get(name,[])+[headers]+[[row[h] for h in headers] for row in rows]
    assert result['Register'][2]==release['native_headers']['Register']
    return result

def validate_native(release):
    native_sheets(release)
    tables=release['tables'];records=tables['Register'];phases=tables['Chronology phases']
    ids=[r['ID'] for r in records];pids=[p['Phase ID'] for p in phases]
    assert len(ids)==len(set(ids)) and len(pids)==len(set(pids))
    assert release['counts']=={'records':len(records),'included':sum(r['Analysis inclusion']=='Included' for r in records),'phases':len(phases)}
    assert set(release['countries_by_record'])==set(ids)
    assert all(isinstance(v,list) and v and all(isinstance(c,str) and c.strip()==c and c for c in v) for v in release['countries_by_record'].values())
    parents={r['ID']:r for r in records}
    for p in phases:
        assert p['Parent record ID'] in parents
        assert bool(p['Parent included (1=yes)'])==(parents[p['Parent record ID']]['Analysis inclusion']=='Included')
        if p['Contested / held (1=yes)']:
            assert p['Minimum date difference (days)'] is None and p['Maximum date difference (days)'] is None
            assert not p['Complete operational (1=yes)'] and not p['Exact complete operational (1=yes)']
        if p['Exact complete operational (1=yes)']:
            assert p['Complete operational (1=yes)'] and p['Evidence category']=='operational_spell'
        if p.get('Phase broad rationales') is not None:
            assert isinstance(p['Phase broad rationales'],list)
            assert set(p['Phase broad rationales'])<=set(tables['Durations'][0]),('Invalid phase rationale',p['Phase ID'])
    assert set(release['source_urls_by_record'])==set(ids)
    assert release['review_counts']['completed_checkpoint_verified_cycles']==50
    return {'records':len(ids),'phases':len(pids),'held_phases':sum(bool(p['Contested / held (1=yes)']) for p in phases)}

def validate_reader(path=READER_PATH, model_path=MODEL_PATH, release=None):
    """Compare every reader cell/link with its checked-in authoring model."""
    from openpyxl import load_workbook
    model=load_json(model_path)
    if release is not None:
        assert model['data_state_sha256']==release['data_state_sha256']
        assert model['counts']==release['counts']
        assert [r[0] for r in model['rows']]==[r['ID'] for r in release['tables']['Register']]
        assert model['source_urls_by_record']==release['source_urls_by_record']
    formulas=load_workbook(path,data_only=False,read_only=False)
    cached=load_workbook(path,data_only=True,read_only=True)
    try:
        assert formulas.sheetnames==cached.sheetnames==['Closures','Methodology']
        assert all(s.sheet_state=='visible' for s in formulas)
        sheet=formulas['Closures'];values=list(cached['Closures'].iter_rows(values_only=True))
        assert sheet.freeze_panes=='B2'
        assert (sheet.max_row,sheet.max_column)==(len(model['rows'])+1,18)
        headers=[c['name'] for c in model['columns']]
        assert [c.value for c in sheet[1]]==headers and len(set(headers))==18
        assert set(sheet.tables)=={'BorderClosures'}
        table=sheet.tables['BorderClosures'];assert table.ref==f'A1:R{len(model["rows"])+1}'
        assert table.autoFilter.ref==table.ref and [c.name for c in table.tableColumns]==headers
        links=0
        for i,row in enumerate(model['rows'],2):
            for j,expected in enumerate(row,1):
                cell=sheet.cell(i,j);actual=cell.value;cache=values[i-1][j-1]
                if 11<=j<=13:
                    index=j-11
                    if index<len(model['links'][i-2]):
                        url=model['links'][i-2][index];escape=lambda v:str(v).replace('"','""')
                        formula=f'=IFERROR(HYPERLINK("{escape(url)}","{escape(expected)}"),"{escape(expected)}")'
                        assert actual==formula and cache==expected,('Reader source hyperlink',i,j)
                        assert url in model['source_urls_by_record'][row[0]];links+=1
                    else: assert actual in (None,'') and cache in (None,'')
                elif j==10 and expected:
                    assert isinstance(actual,dt.datetime) and actual.date().isoformat()==str(expected)[:10]
                    assert cell.number_format=='dd mmm yyyy'
                else:
                    assert (actual if actual is not None else '')==(expected if expected is not None else ''),('Reader cell differs',i,j)
                    assert cell.data_type!='f',('Unexpected reader formula',i,j)
                assert cache not in ('#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#SPILL!')
        assert sum(r[15] for r in model['rows'])==model['counts']['phases']
        assert sum(str(r[17]).startswith('Included') for r in model['rows'])==model['counts']['included']
        method_values=[c.value for row in formulas['Methodology'] for c in row if c.value is not None]
        for value in [model['title'],model['subtitle'],model['data_state_sha256']]+[v for row in model['methods'] for v in row]:
            assert value in method_values,('Methodology text missing or changed',str(value)[:90])
        assert not any(c.data_type=='f' for row in formulas['Methodology'] for c in row)
    finally:
        formulas.close();cached.close()
    with zipfile.ZipFile(path) as package:
        names=package.namelist();assert len(names)==len(set(names))
        for name in names:
            if name.endswith(('.xml','.rels')): ET.fromstring(package.read(name))
        assert len([n for n in names if n.startswith('xl/tables/') and n.endswith('.xml')])==1
        assert len([n for n in names if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n)])==2
    return {'status':'passed','reader_sha256':sha256(path),'data_state_sha256':model['data_state_sha256'],
        'sheets':['Closures','Methodology'],'records':len(model['rows']),'cells_compared':len(model['rows'])*18,
        'source_hyperlinks':links,'native_excel_tested':False}

def read_manifest(path=MANIFEST_PATH):
    manifest=load_json(path)
    assert manifest['schema_version']==1
    for name,item in manifest['files'].items():
        file=(ROOT/item['path']).resolve()
        assert file.is_relative_to(ROOT) and file.is_file(),('Missing release artifact',name)
        assert sha256(file)==item['sha256'],('Release artifact changed after manifest binding',name)
    for file,download in [('data','companion_json'),('reader','reader'),('companion_zip','companion_zip')]:
        assert manifest['files'][file]['sha256']==manifest['downloads'][download]['sha256']
    release=load_json(ROOT/manifest['files']['data']['path'])
    assert release['data_state_sha256']==manifest['data_state_sha256'] and release['counts']==manifest['counts']
    return manifest
