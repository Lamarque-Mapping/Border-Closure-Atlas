#!/usr/bin/env python3
"""Check the prepared workbook against the immutable research snapshot.

This gate checks all original cell payloads/formulas/caches, native parts and
number/date formats, then independently reopens both workbooks read-only to
compare typed cell values. Mutation checks prove that content/formula changes
cannot be accepted as formatting. It never saves through a workbook engine.
"""
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "source/Global_Border_Closures_Register_2016_2026.xlsx"
ORIGINAL_SHA = "f96ad671372e6fb9c1afd564373d954c717f4bf809b75a757b48f54ee08ca024"
spec = importlib.util.spec_from_file_location("presentation", ROOT / "scripts/workbook-presentation.py")
PRESENTATION = importlib.util.module_from_spec(spec)
spec.loader.exec_module(PRESENTATION)


def main():
    original_bytes = subprocess.run(["git", "show", "HEAD:source/Global_Border_Closures_Register_2016_2026.xlsx"],
                                    cwd=ROOT, check=True, capture_output=True).stdout
    assert hashlib.sha256(original_bytes).hexdigest() == ORIGINAL_SHA
    with tempfile.TemporaryDirectory(prefix="atlas-presentation-check-") as directory:
        original = Path(directory) / "original.xlsx"
        repaired = Path(directory) / "repaired.xlsx"
        original.write_bytes(original_bytes)
        subprocess.run([sys.executable, str(ROOT / "scripts/repair-workbook-package.py"), str(original),
                        "--repair", str(repaired)], check=True, capture_output=True)
        report = PRESENTATION.verify_presentation(repaired, WORKBOOK)
        assert report["worksheet_count"] == 55
        assert report["cell_payloads_preserved"] == 363749
        assert report["formula_payloads_preserved"] == 37762
        # Mutation regressions use a real research worksheet, preserving its
        # structure. Only presentation attributes may be removed by the check.
        _, contents, _ = PRESENTATION.read_package(repaired)
        raw = contents["xl/worksheets/sheet1.xml"]
        tree = ET.fromstring(raw)
        cell = tree.find(".//" + PRESENTATION.M + "c[@r='A6']")
        assert cell is not None
        for node in cell.iter():
            if node.text: node.text = "NOT_THE_RESEARCH_VALUE"; break
        assert PRESENTATION.stripped_sheet(raw) != PRESENTATION.stripped_sheet(ET.tostring(tree))
        tree = ET.fromstring(raw)
        formula = tree.find(".//" + PRESENTATION.M + "f")
        assert formula is not None
        formula.text = "1+1"
        assert PRESENTATION.stripped_sheet(raw) != PRESENTATION.stripped_sheet(ET.tostring(tree))
        tree = ET.fromstring(raw)
        cache = tree.find(".//" + PRESENTATION.M + "c[@r='E6']/" + PRESENTATION.M + "v")
        assert cache is not None
        cache.text = "ALTERED_CACHE"
        assert PRESENTATION.stripped_sheet(raw) != PRESENTATION.stripped_sheet(ET.tostring(tree))

        import openpyxl  # Read-only independent parser; never used to author.
        typed_cells = 0
        for data_only in [False, True]:
            before = openpyxl.load_workbook(repaired, read_only=True, data_only=data_only)
            after = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=data_only)
            assert before.sheetnames == after.sheetnames
            for name in before.sheetnames:
                rows_before = before[name].iter_rows()
                rows_after = after[name].iter_rows()
                for row_before, row_after in zip(rows_before, rows_after, strict=True):
                    assert len(row_before) == len(row_after)
                    for left, right in zip(row_before, row_after, strict=True):
                        assert left.value == right.value, f"Typed cell differs: {name} {left.coordinate}"
                        assert left.data_type == right.data_type, f"Cell type differs: {name} {left.coordinate}"
                        assert left.number_format == right.number_format, f"Date/number display differs: {name} {left.coordinate}"
                        typed_cells += int(left.value is not None)
            before.close(); after.close()
        with ZipFile(WORKBOOK) as package:
            styles = ET.fromstring(package.read("xl/styles.xml"))
            style_count = len(styles.find(PRESENTATION.M + "cellXfs"))
            resources = json.loads((ROOT / "source/workbook-presentation.json").read_text())
            for index, name in enumerate(resources["sheets"], 1):
                sheet = ET.fromstring(package.read(f"xl/worksheets/sheet{index}.xml"))
                expected = ET.fromstring(resources["sheets"][name]["sheetViews"])
                assert PRESENTATION.xml(sheet.find(PRESENTATION.M + "sheetViews")) == PRESENTATION.xml(expected)
                assert sheet.find(PRESENTATION.M + "sheetViews")[0].get("showGridLines") == "0"
        report.update({"typed_values_and_cached_results_preserved": True,
                       "typed_nonempty_cell_checks": typed_cells,
                       "date_and_number_formats_preserved": True,
                       "cell_formula_cache_mutation_regressions": 3,
                       "cell_styles": style_count,
                       "native_excel_tested": False})
        report.pop("changed_members")
        print(json.dumps(report, indent=2))


if __name__ == "__main__": main()
