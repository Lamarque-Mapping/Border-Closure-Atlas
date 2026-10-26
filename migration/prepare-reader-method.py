#!/usr/bin/env python3
"""Prepare the reader authoring model and public methods from audited inputs."""
import gzip,json,re,subprocess,sys
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from release_common import load_json
release=load_json(ROOT/'source/current-release.json.gz')
inputs=load_json(ROOT/'migration/audit-inputs.json.gz')
def original(path):return json.loads(subprocess.check_output(['git','-C',str(ROOT),'show',inputs['baseline_commit']+':'+path]))
def save(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n')
method=original('source/release-method.json');model=original('source/current-reader-model.json')
method.update(release_id='2026.10.26-full-audit-2',release_date='2026-10-26',evidence_cutoff='2026-10-26')
urls=release['release_metadata']['distinct_recorded_urls']
method['counts'].update(records=829,included_records=826,phases=1069,distinct_recorded_urls=urls)
method['history'].update(additional_full_register_reviews=2,total_register_wide_reviews=7)
method['full_register_audits']=release['audit_integration']
old=method['full_text']
old=old.replace('The resulting register contains 829 records and 1,010 phases, supported by a catalogue of 4,680 distinct recorded URLs.','At that stage the register contained 829 records and 1,010 phases.')
old=old.replace('The five register-wide reviews covered successive versions of the dataset, while the later targeted audits examined selected records.','Those five register-wide reviews covered successive versions of the dataset; the REC50 targeted audits examined selected records.')
addition=('Two additional full-register source audits then each assessed all 829 entries, using six regional AI reviewers and coordinator adjudication. Regional assignments were rotated for the second pass. Reviewers checked cited evidence and searched for missing or suspicious dates, reopening information, affected movement, exceptions and rationale attribution. The first pass made substantive corrections to 292 records; the second made further substantive corrections to 339 records and presentation-only changes to nine. Across the two passes, 505 distinct records received substantive changes. The passes added 26 and 33 phases respectively, bringing the register to 1,069 phases and 826 records with at least one included phase. The excluded road disruption and included rail suspension within SA025 remain separate. Each record received a fresh readable key-source check in the second pass; some other sources were available only as indexed extracts or remained inaccessible. Completion records an assessment, not verification of every claim. Conflicting dates, unmatched reopenings and limits on continuity remain explicit. These AI source reviews do not constitute independent human or different-model validation. The original study window ends on 21 October 2026; later checks through 26 October retain their actual evidence dates and do not backdate an onset or extrapolate closure.')
method['full_text']=old+'\n\n'+addition+f' The current source catalogue contains {urls:,} distinct recorded URLs, not a count of independent reports or successful full-text inspections.'
method['counting_notes'] += ['Two further whole-register audits each covered all 829 records; their changes overlap, so 292 and 339 must not be added as distinct changed records.','There are 826 included parents, but phase-level exclusions still apply.','The default >24h view requires evidence of a strictly longer restriction; exact numeric duration is a separate, more restrictive analysis subset.']
method['account_source']='Historical programme account supplied by Hugh Lamarque; additional full-register audit counts reconciled to saved findings, adjudications and phase migration.'
save(ROOT/'source/release-method.json',method)
release['release_metadata']['research_method']['full_history']=method['full_text']
release['release_metadata']['research_method']['verification_statement']='Two additional full-register AI source assessments each covered all 829 records. Completed review does not establish every claim; source-access limits and unresolved dates remain explicit. No independent human or different-model verification is claimed.'
release['release_metadata']['research_method']['current_vs_historical']='Current accepted reader and chronology supersede earlier supporting-table assessments; original source history remains available.'
raw=(json.dumps(release,ensure_ascii=False,indent=2)+'\n').encode();(ROOT/'source/current-release.json.gz').write_bytes(gzip.compress(raw,mtime=0))
model.update(rows=inputs['reader']['rows'],counts=release['counts'],data_state_sha256=release['data_state_sha256'],
    source_urls_by_record=release['source_urls_by_record'],current_review_by_record=release['current_review_by_record'],review_counts=release['review_counts'],
    release_method=method,methodology_full_text=method['full_text'],subtitle='829 records. 826 with included phases. 1,069 phases. Audited release: 26 October 2026.')
# Reuse actual hyperlink targets from the adjudicated workbook. Reading with
# openpyxl does not resave or alter this source workbook.
wb=load_workbook(sys.argv[1] if len(sys.argv)>1 else ROOT/'source/Global_Border_Closures_Master_2016_2026.xlsx',data_only=False,read_only=True)
links=[]
for row in wb['Closures'].iter_rows(min_row=2):
    values=[]
    for c in row[10:13]:
        m=re.search(r'HYPERLINK\("((?:[^"]|"")*)"',c.value or '')
        if m:values.append(m.group(1).replace('""','"'))
    links.append(values)
wb.close();model['links']=links
methods=[]
for label,value in model['methods']:
    if label=='Final catalogue':value=f'The current register has 829 records and 1,069 phases; 826 records contain at least one included phase. There are {urls:,} distinct recorded source URLs.'
    elif label=='Review coverage limits':value='The first five register-wide reviews covered successive versions. Two further full-register audits each assessed all 829 current entries. Assessed does not mean all attributes or dates are established.'
    elif label=='Extent of review':value='The earlier 50-cycle programme added 128 records and 183 phases and included 160 full-record source reviews and 176 targeted questions across 162 parents. Subsequent full-register audits covered all 829 entries twice; review populations overlap.'
    elif label=='What the URL count means':value=f'There are {urls:,} distinct recorded HTTP(S) URL strings in the current per-record catalogues. Shared reports, mirrors and repeat retrievals are not independent witnesses. Recorded access can be full text, indexed extracts or unsuccessful.'
    elif label=='Duration units':value='Calendar-date differences are not measured elapsed hours. The default >24h view requires a sourced strictly-over-24-hour assessment or a defensible minimum bound. Exactly 24 hours, shorter restrictions and unestablished durations remain available under All. Policies and announced windows retain their evidence labels; duration charts require completed operational phases with supported exact dates.'
    elif label=='Country and entity filters':value='Country and territory filters use the reviewed named entities. Traffic-specific qualifications remain explicit, including railway-only Kaliningrad scope in EU007. Corrected country assignments do not create coordinates or imply closure at every crossing.'
    elif label=='Citation':value=value.replace('25 October 2026','26 October 2026')
    methods.append([label,value])
methods.extend([['Additional full-register audits',addition],['Changes across both passes','505 distinct entries received substantive changes; the second pass also contains nine presentation-only changes. The register gained 59 phases. Unchanged records were also assessed and have source-access records.'],['Mixed inclusion','SA025 contributes only its documented Benapole rail-freight suspension. Its road disruption remains excluded. GL18, SUP-004 and R3-SC-004 remain wholly outside the map and analysis.'],['Later evidence','The original study period ends on 21 October 2026. Source checks and status observations through 26 October retain their actual dates. Neither a review date nor absence of a reopening report is evidence of continuous closure.']])
model['methods']=methods
save(ROOT/'source/current-reader-model.json',model)
print(json.dumps({'counts':model['counts'],'links':sum(map(len,links)),'methods':len(methods),'urls':urls}))
